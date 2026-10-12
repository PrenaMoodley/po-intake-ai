"""Excel review workbook for the sales admin team.

Totals are live Excel formulas, not pasted values, so anyone can audit the
arithmetic in Excel and the workbook stays correct if they fix a quantity.
A 'Check' column compares the formula result with what Python calculated.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import IntakeResult
from .validate import line_total, money

HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
STATUS_FILL = {
    "ready": PatternFill("solid", fgColor="D9EAD3"),
    "review": PatternFill("solid", fgColor="FFF2CC"),
    "blocked": PatternFill("solid", fgColor="F4CCCC"),
}
ZAR = '"R" #,##0.00'


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill = HEADER, HEADER_FILL
    ws.freeze_panes = "A2"
    return ws


def write_workbook(results: list[IntakeResult], path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    orders = _sheet(
        wb,
        "Orders",
        [
            "Source",
            "PO number",
            "Customer",
            "Account",
            "Order date",
            "Delivery date",
            "Subtotal (lines)",
            "Subtotal (PO)",
            "VAT (PO)",
            "Total (PO)",
            "Status",
            "Reviewed by",
        ],
        [22, 14, 30, 10, 12, 13, 15, 15, 12, 15, 10, 14],
    )
    lines = _sheet(
        wb,
        "Lines",
        [
            "PO number",
            "Line",
            "Customer ref",
            "SKU",
            "Description",
            "Qty",
            "Unit",
            "Unit price",
            "Line total",
            "Python total",
            "Check",
        ],
        [14, 6, 14, 14, 40, 8, 8, 12, 13, 13, 8],
    )
    issues = _sheet(
        wb,
        "Issues",
        ["PO number", "Line", "Severity", "Code", "Message"],
        [14, 6, 10, 20, 80],
    )

    reviewed = DataValidation(type="list", formula1='"Pending,Approved,Rejected"', allow_blank=True)
    orders.add_data_validation(reviewed)

    for result in results:
        o = result.order
        first_line = lines.max_row + 1
        for line in o.lines:
            r = lines.max_row + 1
            lines.append(
                [
                    o.po_number,
                    line.line_no,
                    line.customer_ref,
                    line.sku,
                    line.description,
                    line.quantity,
                    line.unit,
                    money(line.unit_price),
                    f'=IF(H{r}="","",ROUND(F{r}*H{r},2))',
                    line_total(line.quantity, line.unit_price),
                    f'=IF(H{r}="","",IF(I{r}=J{r},"OK","DIFF"))',
                ]
            )
            for col in ("H", "I", "J"):
                lines[f"{col}{r}"].number_format = ZAR
        last_line = lines.max_row

        r = orders.max_row + 1
        subtotal_formula = f"=SUM(Lines!I{first_line}:I{last_line})" if o.lines else "=0"
        orders.append(
            [
                result.source,
                o.po_number,
                o.customer_name,
                o.account_code,
                o.order_date,
                o.delivery_date,
                subtotal_formula,
                money(o.subtotal_excl_vat),
                money(o.vat_amount),
                money(o.total_incl_vat),
                result.status,
                "Pending",
            ]
        )
        for col in ("G", "H", "I", "J"):
            orders[f"{col}{r}"].number_format = ZAR
        orders[f"K{r}"].fill = STATUS_FILL[result.status]
        reviewed.add(f"L{r}")

        for issue in result.issues:
            issues.append([o.po_number, issue.line_no, issue.severity, issue.code, issue.message])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
